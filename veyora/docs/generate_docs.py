"""Generate VEYORA submission documents (DOCX + PPTX).

Run:  python docs/generate_docs.py
Outputs: docs/out/SRS.docx, DDS.docx, Final_Report.docx, VEYORA_Presentation.pptx

The DOCX files use genuine Word heading styles; the final report contains a
real TOC field (update it in Word after opening). Markdown tables are
rendered as Word tables. Content comes from the Markdown sources in docs/.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

DOCS = Path(__file__).resolve().parent
OUT = DOCS / "out"
OUT.mkdir(exist_ok=True)

INK = RGBColor(0x18, 0x25, 0x21)
MUTED = RGBColor(0x66, 0x75, 0x6F)
ACCENT = RGBColor(0x09, 0x6C, 0x60)


def style_doc(doc: Document, title: str):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.font.color.rgb = INK
    for name, size, color in [("Heading 1", 16, ACCENT), ("Heading 2", 13, INK),
                              ("Heading 3", 11.5, INK)]:
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
    core = doc.core_properties
    core.title = title
    core.author = "Mange Krisha Ashok (TYCS_68)"


def add_inline(paragraph, text: str):
    """Handle **bold**, *italic*, `code` inline (basic)."""
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"; run.font.size = Pt(10)
            run.font.color.rgb = ACCENT
        else:
            run = paragraph.add_run(token[1:-1]); run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_toc_field(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = ("Table of contents — select this text and press F9 (or "
                 "References → Update Table) in Word to populate it.")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(sep)
    run._r.append(hint); run._r.append(end)


def md_to_docx(md_path: Path, out_path: Path, title: str, toc_at: str | None = None):
    doc = Document()
    style_doc(doc, title)

    # Cover block
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VEYORA"); r.font.size = Pt(28); r.font.bold = True
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.font.size = Pt(15); r.font.bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Mange Krisha Ashok · TYCS_68 · 125BTCS2004 · "
                  "Shah & Anchor Kutchhi Engineering College")
    r.font.color.rgb = MUTED
    doc.add_page_break()

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue

        if line.startswith("# "):  # document title -> skip (covered above)
            i += 1; continue

        if line.startswith("## "):
            heading = line[3:].strip()
            doc.add_heading(heading, level=1)
            bare = re.sub(r"^\d+\.\s*", "", heading)
            if toc_at and bare.lower() == toc_at.lower():
                add_toc_field(doc)
                doc.add_page_break()
            i += 1; continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2); i += 1; continue

        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3); i += 1; continue

        # fenced code blocks
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            cp = doc.add_paragraph()
            run = cp.add_run("\n".join(code))
            run.font.name = "Consolas"; run.font.size = Pt(9.5)
            cp.paragraph_format.left_indent = Inches(0.3)
            continue

        # tables
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1" if False else "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    tc = table.cell(ri, ci)
                    tc.text = ""
                    add_inline(tc.paragraphs[0], cell)
                    if ri == 0:
                        for run in tc.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        # bullet lists
        if re.match(r"^\s*[-*] ", line):
            while i < len(lines) and re.match(r"^\s*[-*] ", lines[i]):
                item = re.sub(r"^\s*[-*] ", "", lines[i]).rstrip()
                bp = doc.add_paragraph(style="List Bullet")
                add_inline(bp, item)
                i += 1
            continue

        # numbered lists
        if re.match(r"^\s*\d+\. ", line):
            while i < len(lines) and re.match(r"^\s*\d+\. ", lines[i]):
                item = re.sub(r"^\s*\d+\. ", "", lines[i]).rstrip()
                np_ = doc.add_paragraph(style="List Number")
                add_inline(np_, item)
                i += 1
            continue

        # blockquote
        if line.startswith("> "):
            qp = doc.add_paragraph()
            qp.paragraph_format.left_indent = Inches(0.3)
            add_inline(qp, line[2:])
            for run in qp.runs:
                run.font.color.rgb = MUTED; run.italic = True
            i += 1; continue

        # ordinary paragraph
        para = doc.add_paragraph()
        add_inline(para, line.strip())
        i += 1

    doc.save(out_path)
    print(f"wrote {out_path}")


# ---------------------------------------------------------------------------
# Presentation
# ---------------------------------------------------------------------------

def build_pptx():
    from pptx import Presentation
    from pptx.util import Inches as In, Pt as PPt
    from pptx.dml.color import RGBColor as RC

    BG = RC(0xF7, 0xF8, 0xF5)
    HERO = RC(0x10, 0x1C, 0x19)
    INKC = RC(0x18, 0x25, 0x21)
    MUT = RC(0x66, 0x75, 0x6F)
    ACC = RC(0x09, 0x6C, 0x60)
    SAGE = RC(0xE8, 0xF1, 0xEC)

    prs = Presentation()
    prs.slide_width = In(13.333)
    prs.slide_height = In(7.5)
    blank = prs.slide_layouts[6]

    def add_slide(bg=BG):
        slide = prs.slides.add_slide(blank)
        fill = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        fill.fill.solid(); fill.fill.fore_color.rgb = bg
        fill.line.fill.background()
        return slide

    def text_box(slide, l, t, w, h, text, size=18, color=INKC, bold=False,
                 align=None):
        box = slide.shapes.add_textbox(In(l), In(t), In(w), In(h))
        tf = box.text_frame; tf.word_wrap = True
        first = True
        for chunk in text.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run(); r.text = chunk
            r.font.size = PPt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = "Calibri"
            if align:
                p.alignment = align
        return box

    def bullets(slide, l, t, w, h, items, size=17, color=INKC):
        box = slide.shapes.add_textbox(In(l), In(t), In(w), In(h))
        tf = box.text_frame; tf.word_wrap = True
        first = True
        for item in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run(); r.text = "•  " + item
            r.font.size = PPt(size); r.font.color.rgb = color
            r.font.name = "Calibri"
            p.space_after = PPt(8)

    def accent_bar(slide):
        bar = slide.shapes.add_shape(1, In(0.55), In(0.62), In(0.9), In(0.06))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACC
        bar.line.fill.background()

    def title_of(slide, text):
        text_box(slide, 0.5, 0.35, 12, 0.8, text, size=28, color=INKC, bold=True)
        accent_bar(slide)

    # 1 — Title
    s = add_slide(HERO)
    text_box(s, 0.9, 1.4, 11.5, 1.2, "VEYORA", size=54, color=SAGE, bold=True)
    text_box(s, 0.9, 2.5, 11.5, 1.6,
             "Risk-Aware Project Plan for Development of a\n"
             "Smart Healthcare Application", size=26, color=RC(0xF7, 0xF8, 0xF5))
    text_box(s, 0.9, 4.3, 11.5, 0.5, "“A little less complexity. A lot more care.”",
             size=18, color=ACC)
    text_box(s, 0.9, 5.2, 11.5, 1.4,
             "Mange Krisha Ashok · Roll TYCS_68 · PRN 125BTCS2004\n"
             "Shah & Anchor Kutchhi Engineering College · Software Engineering Practices\n"
             "Individual submission · SDG 3 — Good Health and Well-being",
             size=15, color=MUT)

    # 2 — Problem & scope
    s = add_slide()
    title_of(s, "Problem & Scope")
    bullets(s, 0.7, 1.5, 6.0, 5.2, [
        "Small clinics face sprawling systems or ad-hoc spreadsheets",
        "Assignment demands a working prototype AND a risk-aware plan",
        "In scope: auth + roles, in-clinic appointments, health logs,",
        "in-app reminders, risk register, plan & verification views",
        "Out of scope (explicit): video consults, prescriptions, payments,",
        "SMS/email/push, AI diagnosis, offline sync, hospital integrations",
    ])
    text_box(s, 7.2, 1.6, 5.3, 4.6,
             "Three roles\n\n• Patient — books, cancels, keeps private logs\n"
             "• Doctor — manages assigned appointments\n"
             "• Project manager — risks, plan, tasks, verification",
             size=16, color=INKC)

    # 3 — SDG 3
    s = add_slide()
    title_of(s, "SDG 3 — Good Health and Well-being")
    bullets(s, 0.7, 1.5, 11.9, 4.6, [
        "Less friction for basic care coordination: booking a visit, remembering it,",
        "and keeping simple personal readings in one private place",
        "No clinical claims: validation limits are technical constraints, never diagnosis",
        "Honest scope — features requiring clinical certification are excluded, and the",
        "UI never implies they exist",
    ], size=18)

    # 4 — Design & methodology
    s = add_slide()
    title_of(s, "Design & Methodology")
    bullets(s, 0.7, 1.4, 6.0, 5.4, [
        "React 18 + Vite + TypeScript + Motion (reference design preserved:",
        "ivory/teal palette, Manrope, rounded cards, floating phones)",
        "Flask + SQLAlchemy + SQLite; same-origin /api proxy",
        "Diagrams: use-case, architecture, ER, booking activity & sequence",
        "Fixed implementation order; JIRA backlog (5 epics, 16 issues)",
    ])
    bullets(s, 7.0, 1.4, 5.7, 5.4, [
        "Key design decisions:",
        "• Duplicate-booking safety in the DB (partial unique index)",
        "• Reminders derived live — never stored, never stale",
        "• Role isolation by query shape, not UI hiding",
        "• Estimated vs actual effort kept structurally separate",
    ])

    # 5 — Risk-aware plan
    s = add_slide()
    title_of(s, "Risk-Aware Project Plan")
    rows = [
        ("RISK-01", "Unauthorized access to patient data", "3×5 = 15 High", "Sessions, RBAC, CSRF — Implemented"),
        ("RISK-02", "Duplicate appointment booking", "3×4 = 12 High", "Partial unique index + transaction — Implemented"),
        ("RISK-03", "Invalid or misleading data entry", "3×3 = 9 Mod.", "Dual validation — Implemented"),
        ("RISK-04", "Database loss", "2×4 = 8 Mod.", "File DB + backup command — Planned"),
        ("RISK-05", "Development delay (solo student)", "4×3 = 12 High", "Small scope, buffer, earlier deadline"),
        ("RISK-06", "Stale appointment reminders", "2×2 = 4 Low", "Live computation — Implemented"),
    ]
    table = s.shapes.add_table(len(rows) + 1, 4, In(0.7), In(1.5),
                               In(12.0), In(4.6)).table
    widths = [In(1.5), In(4.6), In(2.1), In(3.8)]
    for ci, w in enumerate(widths):
        table.columns[ci].width = w
    hdr = ["Risk", "Description", "Score (P×I)", "Control"]
    for ci, htxt in enumerate(hdr):
        cell = table.cell(0, ci)
        cell.text = htxt
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.size = PPt(14)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = PPt(13)
    text_box(s, 0.7, 6.3, 12, 0.7,
             "Scores are illustrative planning estimates — project-risk priorities, "
             "not clinical classifications. Owner: Krisha Mange.", size=13, color=MUT)

    # 6 — Implementation
    s = add_slide()
    title_of(s, "Implementation")
    bullets(s, 0.7, 1.5, 6.1, 5.2, [
        "4-step booking: doctor → date → slot → review, with friendly 409",
        "conflict handling that keeps doctor + date selected",
        "Health logs: heart rate, temperature, BP (split systolic/diastolic)",
        "Doctor: complete/cancel with validated terminal transitions",
        "PM workspace: register, 5×5 heatmap, plan, tasks, verification",
    ])
    bullets(s, 7.1, 1.5, 5.6, 5.2, [
        "Security: hashed passwords, HttpOnly SameSite cookies,",
        "CSRF on state changes, per-endpoint role guards",
        "Slots generated from the runtime date (IST) — no hardcoded",
        "schedules; UTC storage, IST display",
        "Idempotent seed; explicit dev-only reset",
    ])

    # 7 — Actual test results
    s = add_slide()
    title_of(s, "Actual Testing Results (2026-10-04)")
    bullets(s, 0.7, 1.5, 12, 5.0, [
        "Backend: 42 pytest tests passed in 31.4 s — log: docs/evidence/pytest-run-2026-10-04.txt",
        "Concurrency: two authenticated clients raced for one slot → exactly one 201 and one 409",
        "Frontend production build (strict TypeScript) — success; no CDN React / runtime Babel",
        "Live smoke test via proxy: login 200, booking 201, duplicate 409, no-CSRF 403, unauth 401",
        "Verification table shows real run dates + evidence paths; nothing recorded before running",
    ], size=17)
    text_box(s, 0.7, 6.1, 12, 0.8,
             "Not claimed: JIRA screenshots, demo video, latency numbers — "
             "these are pending student actions.", size=13, color=MUT)

    # 8 — Analysis & reflection
    s = add_slide()
    title_of(s, "Analysis & Reflection")
    bullets(s, 0.7, 1.5, 6.1, 5.2, [
        "Risk-first design: hardest requirements became DB-level",
        "controls with executable tests",
        "Traceability (REQ → risk → control → test → evidence)",
        "kept every claim verifiable",
        "Honesty as a feature: PENDING until actually run",
    ])
    bullets(s, 7.1, 1.5, 5.6, 5.2, [
        "Limitations acknowledged: dev-server scale,",
        "convention-based timezones, manual browser checklist",
        "Personal reflection (400–500 words) written by the",
        "student from her own experience — guide in docs/",
    ])

    # 9 — Conclusion & future scope
    s = add_slide()
    title_of(s, "Conclusion & Future Scope")
    bullets(s, 0.7, 1.5, 12, 4.4, [
        "A working, deliberately small healthcare prototype delivered through a",
        "documented risk-aware plan — with real evidence at every step",
        "Future (not built, not claimed): rescheduling, doctor availability editing,",
        "approved notification channels, PostgreSQL path, clinical features only",
        "under proper governance",
    ], size=18)

    # 10 — Evidence & demo
    s = add_slide()
    title_of(s, "Evidence of Completion")
    bullets(s, 0.7, 1.5, 12, 4.2, [
        "Test log: docs/evidence/pytest-run-2026-10-04.txt (42 passed)",
        "Build log: docs/evidence/frontend-build-2026-10-04.txt",
        "[PLACEHOLDER] JIRA backlog + board screenshots (student action)",
        "[PLACEHOLDER] Demonstration video link — recorded and uploaded by",
        "Krisha Mange (script in docs/demo-script.md)",
    ], size=18)
    text_box(s, 0.7, 6.0, 12, 0.8,
             "Rubric alignment documented in the final report, section "
             "“Report & Presentation Section”.", size=13, color=MUT)

    path = OUT / "VEYORA_Presentation.pptx"
    prs.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    md_to_docx(DOCS / "srs.md", OUT / "SRS.docx",
               "Software Requirements Specification — VEYORA")
    md_to_docx(DOCS / "dds.md", OUT / "DDS.docx",
               "Detailed Design Specification — VEYORA")
    md_to_docx(DOCS / "final-report.md", OUT / "Final_Report.docx",
               "Risk-Aware Project Plan for Development of a Smart Healthcare Application",
               toc_at="Table of Contents")
    build_pptx()
